"""
Markdown → DOCX из того же источника, что MD и PDF: три файла обязаны
совпадать по содержанию, поэтому конвертер один, а не три ручных вёрстки.
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

INK = RGBColor(0x25, 0x25, 0x4F)
INDIGO = RGBColor(0x3F, 0x3F, 0x8A)
MUTED = RGBColor(0x6B, 0x72, 0x80)


def shade(cell, color: str):
    el = OxmlElement('w:shd')
    el.set(qn('w:fill'), color)
    cell._tc.get_or_add_tcPr().append(el)


def add_runs(paragraph, text: str):
    """Разбираем **жирный**, *курсив* и `код` — остальная разметка в документе не используется."""
    for part in re.split(r'(\*\*[^*]+\*\*|(?<!\w)\*[^*]+\*(?!\w)|`[^`]+`)', text):
        if not part:
            continue
        run = paragraph.add_run()
        if part.startswith('**') and part.endswith('**'):
            run.text, run.bold = part[2:-2], True
        elif part.startswith('`') and part.endswith('`'):
            run.text = part[1:-1]
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
        elif part.startswith('*') and part.endswith('*'):
            run.text, run.italic = part[1:-1], True
        else:
            run.text = part


def build(md: str, dst: str):
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(1.8)
        section.left_margin = section.right_margin = Cm(1.7)

    lines = md.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i]

        if line.startswith('```'):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                block.append(lines[i])
                i += 1
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.5)
            run = para.add_run('\n'.join(block))
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            i += 1
            continue

        if line.startswith('|') and i + 1 < len(lines) and re.match(r'^\|[\s:|-]+\|$', lines[i + 1]):
            head = [c.strip() for c in line.strip('|').split('|')]
            i += 2
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                rows.append([c.strip() for c in lines[i].strip('|').split('|')])
                i += 1
            table = doc.add_table(rows=1, cols=len(head))
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for cell, text in zip(table.rows[0].cells, head):
                cell.text = ''
                add_runs(cell.paragraphs[0], text)
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.size = Pt(9)
                    run.font.color.rgb = INK
                shade(cell, 'EEF0F6')
            for row in rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = ''
                    add_runs(cell.paragraphs[0], text)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9)
            doc.add_paragraph()
            continue

        if re.match(r'^#{1,4} ', line):
            level = len(line) - len(line.lstrip('#'))
            text = line[level + 1:]
            if level == 1:
                para = doc.add_paragraph()
                run = para.add_run(text)
                run.bold = True
                run.font.size = Pt(22)
                run.font.color.rgb = INK
            else:
                # add_heading без текста создаёт абзац без ранов: писать в
                # para.runs[-1] нельзя, раны нужно добавить самому.
                para = doc.add_heading('', level=min(level, 4))
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = para.add_run(text)
                run.bold = True
                run.font.color.rgb = INK if level == 2 else INDIGO
                run.font.size = Pt(15 if level == 2 else 12)
            i += 1
            continue

        if line.strip() in ('---', '***'):
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(6)
            border = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:color'), 'C9CDDB')
            border.append(bottom)
            para._p.get_or_add_pPr().append(border)
            i += 1
            continue

        if re.match(r'^[-*] ', line) or re.match(r'^\d+\. ', line):
            ordered = bool(re.match(r'^\d+\. ', line))
            while i < len(lines) and (re.match(r'^[-*] ', lines[i]) or re.match(r'^\d+\. ', lines[i])):
                text = re.sub(r'^([-*]|\d+\.) ', '', lines[i])
                para = doc.add_paragraph(style='List Number' if ordered else 'List Bullet')
                add_runs(para, text)
                i += 1
            continue

        if line.startswith('> '):
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.6)
            add_runs(para, line[2:])
            for run in para.runs:
                run.italic = True
                run.font.color.rgb = MUTED
            i += 1
            continue

        if line.strip():
            add_runs(doc.add_paragraph(), line)
        i += 1

    doc.save(dst)
    print(f'{dst} готов')


if __name__ == '__main__':
    with open(sys.argv[1], encoding='utf-8') as f:
        build(f.read(), sys.argv[2])
